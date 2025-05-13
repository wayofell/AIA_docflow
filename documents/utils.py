import os
import PyPDF2
from PIL import Image
import pytesseract
from pdfminer.high_level import extract_text
import io
import logging
import chardet
import docx
import openpyxl
from xml.etree import ElementTree as ET

# Настройка логирования
logger = logging.getLogger(__name__)

def extract_text_from_pdf(file_path):
    """
    Extract text from a PDF file using a combination of PyPDF2 and pdfminer.six
    """
    if not os.path.exists(file_path):
        logger.error(f"PDF file not found: {file_path}")
        return ""
        
    # First try with PyPDF2
    text = ""
    try:
        with open(file_path, "rb") as pdf_file:
            pdf_reader = PyPDF2.PdfReader(pdf_file)
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        logger.info(f"PyPDF2 extracted {len(text)} characters from {file_path}")
    except Exception as e:
        logger.error(f"PyPDF2 extraction failed: {e}")
    
    # If PyPDF2 didn't get much text, try with pdfminer
    if len(text.strip()) < 50:
        try:
            text = extract_text(file_path)
            logger.info(f"pdfminer extracted {len(text)} characters from {file_path}")
        except Exception as e:
            logger.error(f"pdfminer extraction failed: {e}")
    
    return text

def extract_text_from_image(file_path):
    """
    Extract text from image files (JPG, JPEG, PNG, GIF) using pytesseract
    """
    if not os.path.exists(file_path):
        logger.error(f"Image file not found: {file_path}")
        return ""
        
    try:
        image = Image.open(file_path)
        text = pytesseract.image_to_string(image)
        logger.info(f"pytesseract extracted {len(text)} characters from {file_path}")
        return text
    except Exception as e:
        logger.error(f"Image text extraction failed: {e}")
        return ""

def extract_text_from_docx(file_path):
    """
    Extract text from DOCX files
    """
    if not os.path.exists(file_path):
        logger.error(f"DOCX file not found: {file_path}")
        return ""
    
    try:
        doc = docx.Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        logger.info(f"docx extracted {len(text)} characters from {file_path}")
        return text
    except Exception as e:
        logger.error(f"DOCX text extraction failed: {e}")
        return ""

def extract_text_from_xlsx(file_path):
    """
    Extract text from XLSX files
    """
    if not os.path.exists(file_path):
        logger.error(f"XLSX file not found: {file_path}")
        return ""
    
    try:
        workbook = openpyxl.load_workbook(file_path, read_only=True)
        text = []
        for sheet in workbook.worksheets:
            for row in sheet.rows:
                row_text = [str(cell.value) if cell.value is not None else "" for cell in row]
                text.append(" ".join(row_text))
        
        result = "\n".join(text)
        logger.info(f"xlsx extracted {len(result)} characters from {file_path}")
        return result
    except Exception as e:
        logger.error(f"XLSX text extraction failed: {e}")
        return ""

def extract_text_from_text_file(file_path):
    """
    Extract text from TXT, MD files
    """
    if not os.path.exists(file_path):
        logger.error(f"Text file not found: {file_path}")
        return ""
    
    try:
        # Detect encoding
        with open(file_path, 'rb') as file:
            raw_data = file.read()
            result = chardet.detect(raw_data)
            encoding = result['encoding']
        
        # Read the file with detected encoding
        with open(file_path, 'r', encoding=encoding) as file:
            text = file.read()
        
        logger.info(f"Text file extracted {len(text)} characters from {file_path}")
        return text
    except Exception as e:
        logger.error(f"Text file extraction failed: {e}")
        return ""

def extract_text_from_svg(file_path):
    """
    Extract text from SVG files
    """
    if not os.path.exists(file_path):
        logger.error(f"SVG file not found: {file_path}")
        return ""
    
    try:
        tree = ET.parse(file_path)
        root = tree.getroot()
        text_elements = root.findall(".//{http://www.w3.org/2000/svg}text")
        text = "\n".join([elem.text for elem in text_elements if elem.text])
        logger.info(f"SVG extracted {len(text)} characters from {file_path}")
        return text
    except Exception as e:
        logger.error(f"SVG text extraction failed: {e}")
        return ""

def extract_text_from_file(file_path):
    """
    Extract text from a file based on its extension
    """
    if not file_path or not os.path.exists(file_path):
        logger.error(f"File not found or invalid path: {file_path}")
        return ""
        
    _, file_extension = os.path.splitext(file_path)
    file_extension = file_extension.lower().replace('.', '')
    
    logger.info(f"Extracting text from file: {file_path} with extension: {file_extension}")
    
    if file_extension == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_extension in ['jpg', 'jpeg', 'png', 'gif', 'heic']:
        return extract_text_from_image(file_path)
    elif file_extension in ['docx', 'doc']:
        return extract_text_from_docx(file_path)
    elif file_extension in ['xlsx', 'xls']:
        return extract_text_from_xlsx(file_path)
    elif file_extension in ['txt', 'md']:
        return extract_text_from_text_file(file_path)
    elif file_extension == 'svg':
        return extract_text_from_svg(file_path)
    else:
        logger.warning(f"Unsupported file extension for text extraction: {file_extension}")
        return "" 